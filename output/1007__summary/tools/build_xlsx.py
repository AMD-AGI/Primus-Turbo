#!/usr/bin/env python3
"""Build rounds.xlsx (Chinese) and rounds.en.xlsx (English) from the round ledgers.

CPU only. Inputs (paths relative to output/1007__summary):
  parts/rounds.csv          Chinese ledger (UTF-8 BOM)
  parts/en/rounds.en.csv    English ledger (same rows and columns)
  parts/c_rounds.md, parts/en/c_rounds.md   epoch table of section c.1 (legend sheet)
Both ledgers carry the extra columns stage, machine_short, improved (1/0), kind, gain.

Usage: python3 -I tools/build_xlsx.py
"""
import csv
import math
import os
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

IMPROVING_KINDS = {'accepted', 'promoted', 'new-default', 'hand-champion', 'e2e-gain', 'gemm-fix', 'port'}
BAD_KINDS = {'rejected', 'failed', 'rolled-back'}
GREY_KINDS = {'baseline', 're-measure', 'platform', 'info', 'null', 'retracted'}

KIND_LABEL = {
    'zh': {'baseline': '基线', 'accepted': '接受', 'promoted': '晋升', 'new-default': '新默认',
           'hand-champion': '手工冠军', 'e2e-gain': 'e2e 提升', 'gemm-fix': 'GEMM 修复', 'port': '移植',
           're-measure': '复测', 'platform': '平台变化', 'rejected': '否决', 'null': '无效果',
           'failed': '失败', 'rolled-back': '已回滚', 'retracted': '撤回/作废', 'info': '信息'},
    'en': {'baseline': 'baseline', 'accepted': 'accepted', 'promoted': 'promoted', 'new-default': 'new default',
           'hand-champion': 'hand-campaign champion', 'e2e-gain': 'e2e gain', 'gemm-fix': 'GEMM fix',
           'port': 'port', 're-measure': 're-measure', 'platform': 'platform change', 'rejected': 'rejected',
           'null': 'null', 'failed': 'failed', 'rolled-back': 'rolled back', 'retracted': 'retracted/void',
           'info': 'info'},
}

# (key, width, zh header, en header); key is a CSV column or a derived field.
ALL_COLS = [
    ('seq', 5.5, '序号', '#'),
    ('stage', 19, '阶段', 'Stage'),
    ('date', 10.5, '日期', 'Date'),
    ('direction', 8, '方向', 'Direction'),
    ('round', 22, '轮次', 'Round'),
    ('machine_short', 11, '机器', 'Machine'),
    ('ms', 18, '算子耗时 ms', 'Op time (ms)'),
    ('tflops', 18, 'TF/s', 'TF/s'),
    ('pct_of_asm', 15, '占 ASM', '% of ASM'),
    ('e2e', 26, 'e2e 耗时 / 吞吐', 'e2e step time / throughput'),
    ('accepted', 18, '结果', 'Result'),
    ('kind', 10, '类别', 'Category'),
    ('improved', 6.5, '有提升', 'Improved'),
    ('gain', 30, '提升幅度', 'Gain'),
    ('content', 56, '优化内容', 'What changed'),
    ('platform_note', 44, '备注（驱动/固件/口径）', 'Notes (driver / firmware / convention)'),
    ('ruler', 34, '尺子 / 口径', 'Ruler / convention'),
    ('job_or_campaign', 26, 'Job / 攻关', 'Job / campaign'),
    ('machine', 34, '机器 / 时钟详情', 'Machine / clock details'),
    ('source', 40, '来源', 'Source'),
]
IMP_COLS = ['seq', 'stage', 'date', 'direction', 'round', 'machine_short', 'kind', 'gain', 'ms', 'tflops',
            'pct_of_asm', 'e2e', 'content', 'platform_note', 'source']
FREEZE = {'all': 'F2', 'imp': 'F2'}

SHEET = {'zh': ('全部轮次', '有提升的轮次', '说明'), 'en': ('All rounds', 'Improvement rounds', 'Legend')}
FONT = {'zh': 'Microsoft YaHei', 'en': 'Calibri'}

DIR_FILL = {'fwd': 'DDEBF7', 'bwd': 'FCE4D6', 'fwd+bwd': 'E4DFEC', 'e2e': 'E2EFDA'}
HDR_FILL = {'all': '1F4E78', 'imp': '375623'}
IMP_FILL = 'FFF2CC'
THIN = Side(style='thin', color='BFBFBF')
MED = Side(style='medium', color='404040')

LEGEND = {
    'zh': {
        'title': 'rounds.xlsx 说明',
        'paras': [
            '范围：Llama-3.1-8B 训练用 attention，算子形状 b4 s8192 hq32 hkv8 d128 bf16，causal，BSHD，GQA 4；硬件 AMD MI455X（gfx1250）。'
            '完整报告见同目录 SUMMARY.md / REPORT-1007.html（c 节是本表的汇总版）。',
            '每行是一个轮次或里程碑，按时间排序；“序号”是本表自己的编号，与 SUMMARY.md c.2 的行号不同（c.2 把连续被否决的轮次合并成一行）。',
            'FLOP 约定：fwd 2.199292e12、bwd 5.498229e12（单方向，op-evolve tools/op_flops.py）。阶段 1–2 的 TF/s 多为 fwd+bwd 合计口径（7.697e12 ÷ fwd+bwd 总 ms），单元格里已注明“合计 / fwd / bwd”。合计口径与单方向的数不能直接比。',
            '不同机器、固件阶段、尺子下的绝对数不能直接比，只比同一阶段、同一尺子内的比值（见 SUMMARY.md c.1 和附录 A）。机器代号见下表。',
            '“有提升”：这一行让所在赛道（Triton 路径、默认 attention 路径、FlyDSL bwd、FlyDSL fwd、e2e 吞吐、产品分支）的最好成绩在本行自己的测量阶段内超过了此前的最好成绩。'
            '基线、复测、固件/驱动带来的提速、被否决/失败/事后回滚的轮次、已撤回的数字不算提升。“提升幅度”写明前后对比和参照对象。第 2 个 sheet 只列这些行。',
            '颜色：黄色底 = 有提升的行；“方向”列 蓝 = fwd、橙 = bwd、紫 = fwd+bwd、绿 = e2e；“结果/类别”列 绿字 = 带来提升，红字 = 否决/失败/回滚，灰字 = 基线/复测/平台变化/信息/撤回。阶段之间用粗线分隔。表头可筛选，前 5 列和表头已冻结。',
            '数据源：parts/rounds.csv（中文）、parts/en/rounds.en.csv（英文）；由 tools/build_xlsx.py 生成（tools/build.sh 一并重建报告和 HTML）。未运行任何 GPU 任务。',
        ],
        'epoch_title': '机器代号（与 SUMMARY.md c.1 相同）',
        'cols_title': '各列含义',
        'cols': [
            ('序号', '本表的时间顺序编号'), ('阶段', 'SUMMARY.md c.2 的 11 个阶段'), ('日期', 'MM-DD，UTC'),
            ('方向', 'fwd / bwd / fwd+bwd / e2e'), ('轮次', 'op-evolve 轮次（标明机器与编号体系）、手工攻关的 arm、或里程碑'),
            ('机器', '机器与固件阶段代号，见上表'), ('算子耗时 ms / TF/s / 占 ASM', '该行的算子读数；口径见“尺子/口径”和单元格内注明'),
            ('e2e 耗时 / 吞吐', 'Llama-3.1-8B 训练单步耗时与 tok/s（若有）'), ('结果', '原始判定（accepted / rejected 等，含 gain）'),
            ('类别', '归一化的结果类别，便于筛选'), ('有提升 / 提升幅度', '见上面的定义'),
            ('优化内容', '这一轮改了什么'), ('备注', '驱动 / 固件 / 时钟 / 口径等'), ('尺子 / 口径', '计时方法、数据、对照对象'),
            ('Job / 攻关', 'op-evolve job 名或攻关名称'), ('机器 / 时钟详情', '机器全名与当时的时钟'), ('来源', '原始记录的路径或会话'),
        ],
    },
    'en': {
        'title': 'rounds.en.xlsx legend',
        'paras': [
            'Scope: attention for Llama-3.1-8B training, op shape b4 s8192 hq32 hkv8 d128 bf16, causal, BSHD, GQA 4, on AMD MI455X (gfx1250). '
            'Full report: SUMMARY.en.md / REPORT-1007.en.html in the same folder (section c summarizes this table).',
            'Each row is one round or milestone, in time order. "#" is this table\'s own numbering and differs from the row numbers of SUMMARY c.2 (c.2 merges consecutive rejected rounds into one row).',
            'FLOP convention: fwd 2.199292e12, bwd 5.498229e12 per direction (op-evolve tools/op_flops.py). In stages 1-2 most TF/s values use the fwd+bwd total convention (7.697e12 / total fwd+bwd ms); each cell says "total / fwd / bwd". Total and per-direction numbers are not comparable.',
            'Absolute numbers from different machines, firmware epochs or rulers are not comparable; compare only ratios within one epoch and ruler (SUMMARY c.1 and Appendix A). Machine codes are listed below.',
            '"Improved": the row raised the best result of its track (Triton path, default attention path, FlyDSL bwd, FlyDSL fwd, e2e throughput, product branch) above the previous best, within the row\'s own measurement epoch. '
            'Baselines, re-measurements, firmware/driver speedups, rejected/failed/rolled-back rounds and retracted numbers do not count. "Gain" gives before -> after and the reference. Sheet 2 lists only these rows.',
            'Colors: yellow background = improvement row; "Direction": blue = fwd, orange = bwd, purple = fwd+bwd, green = e2e; "Result/Category": green = brought a gain, red = rejected/failed/rolled back, grey = baseline/re-measure/platform/info/retracted. A thick line separates stages. Header row has filters; the header and the first 5 columns are frozen.',
            'Data: parts/rounds.csv (Chinese), parts/en/rounds.en.csv (English); generated by tools/build_xlsx.py (tools/build.sh rebuilds the reports and HTML too). No GPU jobs were run.',
        ],
        'epoch_title': 'Machine codes (same as SUMMARY c.1)',
        'cols_title': 'Columns',
        'cols': [
            ('#', 'time-order number in this table'), ('Stage', 'the 11 stages of SUMMARY c.2'), ('Date', 'MM-DD, UTC'),
            ('Direction', 'fwd / bwd / fwd+bwd / e2e'), ('Round', 'op-evolve round (with machine and numbering), hand-campaign arm, or milestone'),
            ('Machine', 'machine and firmware-epoch code, see the table above'), ('Op time (ms) / TF/s / % of ASM', 'the row\'s op readings; convention in "Ruler / convention" and in the cell'),
            ('e2e step time / throughput', 'Llama-3.1-8B training step time and tok/s, if any'), ('Result', 'original verdict (accepted / rejected ..., with gain)'),
            ('Category', 'normalized verdict for filtering'), ('Improved / Gain', 'see the definition above'),
            ('What changed', 'what the round changed'), ('Notes', 'driver / firmware / clock / convention notes'), ('Ruler / convention', 'timing method, data, reference arm'),
            ('Job / campaign', 'op-evolve job or campaign name'), ('Machine / clock details', 'full machine name and clocks at the time'), ('Source', 'path or session of the original record'),
        ],
    },
}


def clean(s):
    s = (s or '').replace('**', '').replace('`', '')
    return ILLEGAL_CHARACTERS_RE.sub('', s).strip()


def vwidth(s):
    return sum(1.9 if unicodedata.east_asian_width(ch) in 'WF' else 1.0 for ch in s)


def nlines(s, width):
    if not s:
        return 1
    usable = max(width * 1.05 - 1.0, 4.0)
    return sum(max(1, math.ceil(vwidth(p) / usable)) for p in s.split('\n'))


def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:]]


def epoch_table(md_path):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    out = []
    for i, l in enumerate(lines):
        if l.startswith('|') and re.match(r'^\|\s*(代号|Code)', l):
            j = i + 2
            while j < len(lines) and lines[j].startswith('|'):
                cells = [clean(c) for c in re.split(r'(?<!\\)\|', lines[j])[1:-1]]
                out.append(cells)
                j += 1
            head = [clean(c) for c in re.split(r'(?<!\\)\|', l)[1:-1]]
            return head, out
    return None, []


def value(row, key, lang):
    v = row.get(key, '')
    if key == 'kind':
        return KIND_LABEL[lang].get(v, v)
    if key == 'improved':
        return '✓' if v == '1' else ''
    if key == 'seq' and v.isdigit():
        return int(v)
    return clean(v)


def style_header(ws, cols, lang, fill):
    font = Font(name=FONT[lang], bold=True, color='FFFFFF', size=10)
    for c, (key, width, zh, en) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=zh if lang == 'zh' else en)
        cell.font = font
        cell.fill = PatternFill('solid', fgColor=fill)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 34


def write_rows(ws, cols, rows, lang):
    base = Font(name=FONT[lang], size=10)
    prev_stage = None
    for r_i, row in enumerate(rows, 2):
        improved = row.get('improved') == '1'
        kind = row.get('kind', '')
        new_stage = row.get('stage') != prev_stage
        prev_stage = row.get('stage')
        height = 1
        for c, (key, width, _, _) in enumerate(cols, 1):
            v = value(row, key, lang)
            cell = ws.cell(row=r_i, column=c, value=v)
            font = base
            if key in ('accepted', 'kind'):
                if kind in IMPROVING_KINDS and improved:
                    font = Font(name=FONT[lang], size=10, color='2E7D32', bold=(key == 'kind'))
                elif kind in BAD_KINDS:
                    font = Font(name=FONT[lang], size=10, color='C00000')
                elif kind in GREY_KINDS:
                    font = Font(name=FONT[lang], size=10, color='7F7F7F', italic=(kind == 'retracted'))
            elif key == 'improved' and improved:
                font = Font(name=FONT[lang], size=11, color='2E7D32', bold=True)
            elif key in ('seq', 'round', 'gain') and improved:
                font = Font(name=FONT[lang], size=10, bold=True)
            elif key == 'stage':
                font = Font(name=FONT[lang], size=10, bold=new_stage, color='000000' if new_stage else '595959')
            cell.font = font
            centered = key in ('seq', 'date', 'direction', 'machine_short', 'improved', 'kind')
            cell.alignment = Alignment(horizontal='center' if centered else 'left', vertical='top', wrap_text=True)
            if key == 'direction' and row.get('direction') in DIR_FILL:
                cell.fill = PatternFill('solid', fgColor=DIR_FILL[row['direction']])
            elif improved:
                cell.fill = PatternFill('solid', fgColor=IMP_FILL)
            cell.border = Border(left=THIN, right=THIN, bottom=THIN, top=MED if (new_stage and r_i > 2) else THIN)
            height = max(height, nlines(str(v), width))
        ws.row_dimensions[r_i].height = min(14.0 * height + 4, 300)


def finish(ws, ncols, nrows, freeze):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(ncols), max(nrows + 1, 2))
    ws.print_title_rows = '1:1'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.zoomScale = 90


def legend_sheet(ws, lang, epoch_md):
    L = LEGEND[lang]
    f = FONT[lang]
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 46
    ws.column_dimensions['D'].width = 40
    ws.column_dimensions['E'].width = 60
    ws.cell(row=1, column=1, value=L['title']).font = Font(name=f, size=14, bold=True)
    r = 3
    for p in L['paras']:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws.cell(row=r, column=1, value=p)
        c.font = Font(name=f, size=10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 14.0 * nlines(p, 24 + 16 + 46 + 40 + 60) + 6
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=L['epoch_title']).font = Font(name=f, size=11, bold=True)
    r += 1
    head, rows = epoch_table(epoch_md)
    widths = [24, 16, 46, 40, 60]
    if head:
        for c, h in enumerate(head[:5], 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font = Font(name=f, size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        r += 1
        for row in rows:
            h = 1
            for c, v in enumerate(row[:5], 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = Font(name=f, size=10, bold=(c == 1))
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
                h = max(h, nlines(v, widths[c - 1]))
            ws.row_dimensions[r].height = min(14.0 * h + 4, 300)
            r += 1
    r += 1
    ws.cell(row=r, column=1, value=L['cols_title']).font = Font(name=f, size=11, bold=True)
    r += 1
    for name, desc in L['cols']:
        a = ws.cell(row=r, column=1, value=name)
        a.font = Font(name=f, size=10, bold=True)
        a.alignment = Alignment(wrap_text=True, vertical='top')
        a.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        b = ws.cell(row=r, column=2, value=desc)
        b.font = Font(name=f, size=10)
        b.alignment = Alignment(wrap_text=True, vertical='top')
        b.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        r += 1
    ws.sheet_view.showGridLines = False


def build(lang, csv_path, epoch_md, out_path):
    rows = read_csv(csv_path)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET[lang][0]
    style_header(ws, ALL_COLS, lang, HDR_FILL['all'])
    write_rows(ws, ALL_COLS, rows, lang)
    finish(ws, len(ALL_COLS), len(rows), FREEZE['all'])

    imp_cols = [c for k in IMP_COLS for c in ALL_COLS if c[0] == k]
    imp = [r for r in rows if r.get('improved') == '1']
    ws2 = wb.create_sheet(SHEET[lang][1])
    style_header(ws2, imp_cols, lang, HDR_FILL['imp'])
    write_rows(ws2, imp_cols, imp, lang)
    finish(ws2, len(imp_cols), len(imp), FREEZE['imp'])

    legend_sheet(wb.create_sheet(SHEET[lang][2]), lang, epoch_md)
    wb.save(out_path)
    return len(rows), len(imp)


def main():
    for lang, csv_path, md, out in (
            ('zh', 'parts/rounds.csv', 'parts/c_rounds.md', 'rounds.xlsx'),
            ('en', 'parts/en/rounds.en.csv', 'parts/en/c_rounds.md', 'rounds.en.xlsx')):
        p = os.path.join(ROOT, csv_path)
        if not os.path.exists(p):
            print('skip', out, '(missing', csv_path + ')')
            continue
        n, k = build(lang, p, os.path.join(ROOT, md), os.path.join(ROOT, out))
        print('%s: %d rows, %d improvement rows' % (out, n, k))


if __name__ == '__main__':
    main()
